import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Palette
    BG_DARK = RGBColor(18, 24, 38)        # #121826 Deep Industrial Navy
    CARD_BG = RGBColor(28, 36, 56)        # #1C2438 Card surface
    CARD_BORDER = RGBColor(45, 58, 90)    # #2D3A5A Border
    PRIMARY_CYAN = RGBColor(0, 212, 255)  # #00D4FF Accent Cyan
    ACCENT_YELLOW = RGBColor(255, 209, 102) # #FFD166 Inspection Yellow
    TEXT_WHITE = RGBColor(245, 247, 250)  # Main text
    TEXT_MUTED = RGBColor(160, 174, 192)  # Secondary text
    CARD_ACCENT = RGBColor(16, 185, 129)  # #10B981 Success green

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Tag pill
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
        tf = tag_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        if subtitle_text:
            p_sub = tf_title.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, badge=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.2)

        tx_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), height - Inches(0.36))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Title
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY_CYAN

        if badge:
            p_b = tf.add_paragraph()
            p_b.text = badge
            p_b.font.size = Pt(10)
            p_b.font.bold = True
            p_b.font.color.rgb = ACCENT_YELLOW
            p_b.space_after = Pt(6)

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = "• " + item
            p_i.font.size = Pt(12)
            p_i.font.color.rgb = TEXT_WHITE
            p_i.space_before = Pt(4)
            p_i.space_after = Pt(2)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: COVER
    # ═════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent decorative top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(11.733), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_CYAN
    bar.line.color.rgb = PRIMARY_CYAN

    # Main Title text frame
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "LAPORAN KEMAJUAN TUGAS AKHIR"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY_CYAN
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "Sistem Deteksi dan Pencatatan Komponen Elektronika Real-Time"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Berbasis YOLOv8 dan Edge Computing (Raspberry Pi 4)"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_YELLOW
    p2.space_after = Pt(16)

    p3 = tf1.add_paragraph()
    p3.text = "Implementasi Computer Vision pada Konveyor Laboratorium menggunakan Dataset ElectroCom61"
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_MUTED

    # Student metadata card
    m_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.8), Inches(11.733), Inches(1.8))
    m_card.fill.solid()
    m_card.fill.fore_color.rgb = CARD_BG
    m_card.line.color.rgb = CARD_BORDER

    mtx = s1.shapes.add_textbox(Inches(1.1), Inches(5.0), Inches(11.133), Inches(1.4))
    mtf = mtx.text_frame
    mtf.word_wrap = True

    mp1 = mtf.paragraphs[0]
    mp1.text = "Nama Peneliti : Ariful Fikri          |          NIM : 2307110474"
    mp1.font.size = Pt(16)
    mp1.font.bold = True
    mp1.font.color.rgb = TEXT_WHITE
    mp1.space_after = Pt(6)

    mp2 = mtf.add_paragraph()
    mp2.text = "Program Studi : Teknik Informatika   |   Fakultas Teknik, Universitas Riau"
    mp2.font.size = Pt(14)
    mp2.font.color.rgb = PRIMARY_CYAN
    mp2.space_after = Pt(4)

    mp3 = mtf.add_paragraph()
    mp3.text = "Tahun Akademik: 2026"
    mp3.font.size = Pt(13)
    mp3.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: DATASET & TRAINING
    # ═════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Bagian 1: Model Machine Learning", "Dataset & Konfigurasi Pelatihan (Training YOLOv8n)", 
               "Proses pelatihan model deteksi objek untuk pengenalan 61 jenis komponen elektronika")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "1. Karakteristik Dataset", [
                 "Dataset: ElectroCom61 (Kaggle Dataset resmi).",
                 "Total Citra: 2.121 gambar beranotasi bounding box presisi.",
                 "Cakupan Kelas: 61 varian komponen elektronika industri & lab.",
                 "Kategori Komponen: Resistor, Kapasitor (Film/SMD/Elektrolit), Dioda, IC (DIP/SOP), Transistor, Sensor, hingga Modul Driver.",
                 "Pembagian Data: 3-Way Split (Train set 70%, Validation set 20%, Test set 10%)."
             ], badge="61 KELAS KOMPONEN")

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "2. Pemilihan Arsitektur", [
                 "Arsitektur Model: YOLOv8 Nano (yolov8n.pt).",
                 "Ukuran Model Kompak: ~6.2 MB (sangat efisien di memori).",
                 "Alasan Pemilihan: Arsitektur teringan dengan rasio mAP/latensi terbaik untuk target perangkat CPU Edge (Raspberry Pi 4).",
                 "Anchor-Free Detection: Meningkatkan sensitivitas deteksi komponen komponen mikro/kecil pada citra konveyor.",
                 "Backbone: C2f Module untuk ekstraksi fitur yang kaya tanpa komputasi berat."
             ], badge="EDGE AI COMPATIBLE")

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0), 
             "3. Konfigurasi Training", [
                 "Jumlah Epoch: 50 Epochs (konvergen optimal).",
                 "Early Stopping: Aktif (Patience = 10 epoch).",
                 "Batch Size: 8 (menyesuaikan kestabilan gradien).",
                 "Resolusi Input: 640 × 640 piksel.",
                 "Optimizer: Adam Optimizer dengan learning rate lr0 = 0.001.",
                 "Hardware Pelatihan: Laptop GPU NVIDIA GeForce GTX 1650 (4GB VRAM) menggunakan PyTorch CUDA."
             ], badge="OPTIMIZED HYPERPARAMETERS")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: EVALUASI & EXPORT ONNX
    # ═════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Bagian 2: Hasil & Evaluasi Model", "Evaluasi Akurasi & Konversi Model ke ONNX Runtime", 
               "Validasi performa model pada data uji serta optimasi format untuk perangkat edge")

    # Metrics summary cards row
    metrics = [
        ("mAP@0.5 (Val)", "91.64%", "Akurasi deteksi pada IoU 0.50", PRIMARY_CYAN),
        ("mAP@0.5 (Test)", "89.65%", "Evaluasi pada citra data uji baru", ACCENT_YELLOW),
        ("Precision", "87.23%", "Ketepatan identifikasi jenis komponen", CARD_ACCENT),
        ("Recall", "89.74%", "Sensitivitas penemuan objek di frame", PRIMARY_CYAN),
    ]

    for idx, (m_title, m_val, m_desc, m_col) in enumerate(metrics):
        cx = Inches(0.8 + idx * 2.98)
        card_m = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.8), Inches(2.8), Inches(1.6))
        card_m.fill.solid()
        card_m.fill.fore_color.rgb = CARD_BG
        card_m.line.color.rgb = CARD_BORDER

        tx = s3.shapes.add_textbox(cx + Inches(0.15), Inches(1.9), Inches(2.5), Inches(1.4))
        tf = tx.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        p2 = tf.add_paragraph()
        p2.text = m_val
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = m_col

        p3 = tf.add_paragraph()
        p3.text = m_desc
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = TEXT_MUTED

    # Bottom detailed cards
    add_card(s3, Inches(0.8), Inches(3.6), Inches(5.7), Inches(3.3),
             "Analisis Performa Model", [
                 "Skor F1-Score mencapai 88.47%, membuktikan keseimbangan yang kuat antara Presisi dan Recall.",
                 "mAP@0.5:0.95 bernilai 58.82% — membuktikan ketepatan bounding box bahkan pada ambang batas IoU ketat.",
                 "Kategori komponen pasif (Resistor, Kapasitor) terdeteksi stabil tanpa false positive berlebih.",
                 "Model mampu membedakan bentuk fisik komponen modul (mis: Relay Module vs Motor Driver) dengan akurat."
             ], badge="EVALUASI METRIK")

    add_card(s3, Inches(6.8), Inches(3.6), Inches(5.7), Inches(3.3),
             "Konversi Model: PyTorch (.pt) ➔ ONNX Runtime", [
                 "Format Target: Open Neural Network Exchange (best.onnx ~12.3 MB).",
                 "Eliminasi Dependensi Berat: Menghilangkan kebutuhan instalasi PyTorch GPU (~3GB) di OS Raspberry Pi.",
                 "Eksekusi Mandiri: Dijalankan via ONNX Runtime CPUExecutionProvider yang dioptimalkan untuk instruksi CPU ARM.",
                 "Warmup Execution: Diuji 1x inferensi kosong saat boot untuk mengalokasikan memory pool agar frame pertama bebas lag."
             ], badge="CROSS-PLATFORM DEPLOYMENT")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: FITUR 1 - ROI & DEBOUNCING
    # ═════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Fitur Aplikasi 1: Logika Cerdas Inspeksi", "Zona Inspeksi (ROI) & Algoritma Debouncing", 
               "Mekanisme pelacakan posisi dan pencegahan pencatatan data ganda pada konveyor")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "1. Konfigurasi Area ROI", [
                 "Area Zona Inspeksi: Kotak persegi kuning di tengah frame citra.",
                 "Rentang Koordinat Normalisasi: (ROI = 0.25, 0.75, 0.1, 0.9).",
                 "Cakupan Fisik: Mengawasi 50% lebar tengah dan 80% tinggi frame kamera konveyor.",
                 "Filter Fokus Inspeksi: Mencegah komponen di tepi luar konveyor atau objek latar belakang tidak sengaja tercatat ke database.",
                 "Perhitungan Titik Pusat: Titik (cx, cy) dihitung dari rata-rata batas bounding box (x1+x2)/2 dan (y1+y2)/2."
             ], badge="BOUNDING REGION OF INTEREST")

    add_card(s4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "2. Pembeda Visual Frame", [
                 "Objek di DALAM Zona Inspeksi: Ditandai bounding box berwarna cerah, aksen siku di setiap sudut, dan label nama kelas + nilai akurasi.",
                 "Objek di LUAR Zona Inspeksi: Ditandai bounding box abu-abu tipis (150, 150, 150) sebagai status monitoring pasif.",
                 "Tracking Aktif: Selama berada dalam zona, data komponen disimpan di memory buffer active_tracking.",
                 "Pembaruan Nilai Akurasi: Sistem terus memperbarui nilai confidence tertinggi yang pernah terbaca selama komponen melintas."
             ], badge="STATE VISUALIZATION")

    add_card(s4, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "3. Logika Debounce (EXIT_DELAY)", [
                 "Toleransi Waktu Keluar: Parameter EXIT_DELAY = 1.0 detik.",
                 "Alur Pencatatan Database: Ketika komponen meninggalkan zona atau hilang dari frame kamera, timer 1 detik mulai berjalan.",
                 "Trigger Simpan SQLite: Jika komponen tidak terdeteksi lagi setelah > 1.0 detik, data dicatat TEPAT 1 KALI ke tabel deteksi.",
                 "Solusi Masalah Nyata: Mencegah pencatatan ganda ratusan kali jika komponen diam atau berjalan pelan di konveyor.",
                 "Reset Otomatis: Entri active_tracking dihapus agar komponen sejenis berikutnya dapat dicatat kembali."
             ], badge="ANTI-DUPLIKASI 1.0 DETIK")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: FITUR 2 - DEDICATED PI CAMERA
    # ═════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Fitur Aplikasi 2: Abstraksi Sensor Hardware", "Integrasi Dedicated Raspberry Pi Camera & Zero-Latency", 
               "Arsitektur penanganan hardware sensor Pi Camera Rev 1.3 native dan bebas buffer delay")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "1. Dedicated Pi Camera (CSI)", [
                 "Hardware Target: Raspberry Pi Camera Module Rev 1.3.",
                 "Koneksi: Jalur bus Camera Serial Interface (CSI) 15-pin kabel pita langsung ke SoC BCM2711.",
                 "Driver Native: Menggunakan library resmi Picamera2 dan libcamera sub-system.",
                 "Keunggulan CSI: Mengeliminasi beban transfer data USB host controller, mengurangi latency pengambilan gambar.",
                 "Fallback Otomatis: Jika dijalankan di PC/laptop pengujian, otomatis fallback ke webcam 0 tanpa error."
             ], badge="NATIVE CSI BUS PIPELINE")

    add_card(s5, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "2. Warmup Retry Sensor", [
                 "Karakteristik Sensor: Sensor CMOS kamera memerlukan waktu stabilisasi eksposur dan gain pada beberapa frame pertama.",
                 "Warmup Retry Polling: Sistem melakukan percobaan capture hingga 15 kali (1.5 detik) saat inisialisasi awal sensor.",
                 "Pencegahan Frame Korup: Memastikan frame yang masuk ke model AI selalu valid dan tidak kosong (None).",
                 "Suppression Log: Menyembunyikan pesan debug internal V4L2 agar terminal pengujian tetap bersih dan terfokus."
             ], badge="ROBUST HARDWARE INIT")

    add_card(s5, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "3. Threaded Bufferless Capture", [
                 "Masalah Buffer Kamera Standar: Pembacaan synchronous default menyebabkan penumpukan frame lama (antrean delay bertambah jika FPS turun).",
                 "Solusi Dedicated Thread: Class CameraStream menjalankan daemon thread terpisah yang terus menerus menyegarkan frame terbaru.",
                 "Bufferless Snapshot: Thread utama inferensi AI selalu mengambil frame terbaru instan tanpa menunggu hardware kamera.",
                 "Sinkronisasi Mutex: Memastikan tidak terjadi perebutan resource memory (race condition) antar thread."
             ], badge="REAL-TIME ZERO LATENCY")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: FITUR 3 - WEB DASHBOARD & EKSPOR
    # ═════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Fitur Aplikasi 3: Antarmuka & Telemetri", "Web Dashboard Real-Time, Kontrol Sesi & Ekspor Data", 
               "Antarmuka pengguna industrial dengan pemisahan saluran video dan data telemetri")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "1. Dual-Channel Streaming", [
                 "Saluran 1 - Visual Feed: Protokol HTTP MJPEG multipart/x-mixed-replace pada endpoint /video/stream dengan interval 40ms (~25 FPS stabil).",
                 "Saluran 2 - Telemetri AI: Protokol Server-Sent Events (SSE) text/event-stream pada endpoint /results/stream.",
                 "Push Data Real-Time: Data telemetri (waktu inferensi ms, daftar objek, log event, barang baru) di-push otomatis ke browser tanpa polling berulang.",
                 "Ringan di Jaringan: Tidak memerlukan WebSocket berlebih, bekerja optimal di jaringan lokal WiFi laboratorium."
             ], badge="MJPEG + SSE ARCHITECTURE")

    add_card(s6, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "2. Manajemen Sesi Pengujian", [
                 "Tombol 'Mulai Deteksi': Mengaktifkan kamera, memicu thread inferensi YOLO, dan otomatis mengosongkan tabel riwayat sesi sebelumnya.",
                 "Data Bersih Setiap Pengujian: Mencegah data komponen dari batch pengujian lama bercampur dengan pengujian baru.",
                 "Slider Ambang Batas Default 60%: Nilai confidence threshold dipasang optimal di 60% untuk mencegah deteksi palsu (false positive).",
                 "Adaptif Pencahayaan: Penguji tetap dapat menyesuaikan nilai (10%–95%) jika intensitas cahaya lampu konveyor berubah."
             ], badge="DYNAMIC SESSION CONTROL")

    add_card(s6, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "3. Penyimpanan & Ekspor Data", [
                 "Database Lokal: SQLite3 (data/history.db) — menyimpan kolom ID, Waktu (YYYY-MM-DD HH:MM:SS), Kelas Komponen, dan Confidence.",
                 "Ekspor Otomatis: Saat tombol 'Hentikan' ditekan, sistem menampilkan popup konfirmasi untuk mengunduh laporan pengujian.",
                 "Format CSV/Excel: Endpoint /api/history/export menghasilkan file CSV langsung dari memory buffer.",
                 "Encoding utf-8-sig: Dilengkapi Byte Order Mark (BOM) agar tabel terbuka rapi dan terbaca sempurna di Microsoft Excel."
             ], badge="AUTOMATIC REPORT EXPORT")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: FITUR 4 - OPTIMASI RASPBERRY PI 4
    # ═════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Fitur Aplikasi 4: Komputasi Edge", "Optimasi Khusus Perangkat Raspberry Pi 4", 
               "Strategi rekayasa perangkat lunak untuk menjaga kestabilan sistem pada hardware CPU terbatas")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "1. Resolusi Tensor 320×320", [
                 "Parameter: IMGSZ = 320 (Default kode).",
                 "Penurunan Kompleksitas Komputasi: Resolusi 320×320 mengurangi jumlah operasi perkalian matriks hingga 4× lipat dibandingkan 640×640.",
                 "Kecepatan Inferensi: Menurunkan latensi eksekusi dari ~450ms menjadi ~120-150ms per frame pada prosesor ARM Cortex-A72.",
                 "Akurasi Terjaga: Struktur komponen elektronika pada dataset tetap terdeteksi dengan jelas pada skala resolusi ini."
             ], badge="4X INFERENCE SPEEDUP")

    add_card(s7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "2. Mekanisme Frame Skipping", [
                 "Parameter: FRAME_SKIP = 3 (Default kode).",
                 "Prinsip Kerja: Inferensi neural network hanya dieksekusi 1 kali pada setiap 3 frame kamera (~10 kali/detik dari kamera 30 FPS).",
                 "Hemat Siklus CPU: Menghilangkan 66% beban kalkulasi inferensi yang tidak perlu.",
                 "Video Tetap Halus: Pada 2 frame yang dilewati, sistem tetap me-render frame visual dengan koordinat deteksi terakhir.",
                 "Hasil: Tampilan video di browser tetap mulus dan bebas patah-patah."
             ], badge="FRAME_SKIP = 3")

    add_card(s7, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "3. Throttle 10ms & Efisiensi Memori", [
                 "Jeda Mikro CPU: Penerapan time.sleep(0.01) (10 milidetik) pada setiap akhir iterasi loop utama.",
                 "Pencegahan Thermal Throttling: Mencegah CPU core berada pada utilisasi 100% konstan yang menyebabkan panas berlebih (overheat) dan freeze.",
                 "Kualitas JPEG Optimal: Kompresi JPEG_QUALITY = 65 menghemat throughput bandwidth WiFi lokal.",
                 "Bebas Dependensi Raksasa: Menghilangkan dependensi PyTorch GPU sehingga aplikasi aman di MicroSD 16GB."
             ], badge="ANTI-OVERHEATING & STABLE")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: PROGRESS & NEXT STEPS
    # ═════════════════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Bagian 4: Kesimpulan Progres", "Status Capaian Progres & Rencana Pengujian Lanjutan", 
               "Rekapitulasi pekerjaan yang telah diselesaikan serta target pengujian menuju sidang akhir")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0),
             "Capaian yang Telah Diselesaikan (100%)", [
                 "Pelatihan & Verifikasi Model AI: Model YOLOv8n terlatih 50 epoch pada 61 kelas dengan mAP@0.5 mencapai 91.64%.",
                 "Ekspor ONNX Runtime: Model berhasil dioptimalkan ke format ONNX dan berjalan mandiri via CPU tanpa PyTorch GPU.",
                 "Implementasi Logika ROI & Debounce: Logika zona kuning dan delay keluar 1.0 detik sukses mengeliminasi pencatatan data ganda.",
                 "Integrasi Hardware Raspberry Pi: Berhasil mengintegrasikan modul Raspberry Pi Camera Rev 1.3 via Picamera2.",
                 "Pengembangan Web Dashboard: Tampilan dark-mode industrial dengan streaming MJPEG, SSE telemetri, dan ekspor CSV selesai 100%.",
                 "Dokumentasi & Versi Kode: Repositori GitHub telah terorganisir rapi dengan petunjuk instalasi dan referensi ilmiah."
             ], badge="MILESTONES COMPLETED (100%)")

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
             "Rencana Pengujian Lanjutan (Next Steps)", [
                 "Pengujian Variasi Kecepatan Konveyor: Menguji akurasi deteksi model pada kecepatan putaran treadmill lambat, sedang, dan cepat.",
                 "Pengujian Variasi Pencahayaan: Menguji ketahanan akurasi pada berbagai intensitas cahaya lux lampu laboratorium.",
                 "Uji Ketahanan Durasi Operasional (Stress-Test): Menjalankan sistem selama 1-2 jam kontinu untuk memantau stabilitas suhu CPU Raspberry Pi 4.",
                 "Penyusunan Naskah Skripsi: Melengkapi penulisan Bab 4 (Hasil dan Pembahasan) serta Bab 5 (Kesimpulan dan Saran) berdasarkan data pengujian lapangan.",
                 "Persiapan Sidang Akhir: Pembuatan video demo operasional sistem pada konveyor laboratorium."
             ], badge="ACTION PLAN SKU")

    # Save presentation
    output_path = "Presentasi_Progress_Skripsi.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
